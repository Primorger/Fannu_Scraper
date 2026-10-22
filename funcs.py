#Requirements: openpyxl, requests

import json
from collections.abc import Mapping
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlsplit

import requests
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


API_BASE_URL = (
    "https://n2kbackbone.eea.europa.eu/"
    "n2kbackbone_back/api/PublicData"
)
INPUT_PATH = Path(__file__).resolve().parent / "input" / "input.json"
OUTPUT_PATH = Path(__file__).resolve().parent / "output" / "output.xlsx"
HEADERS = (
    "Zone",
    "Group",
    "Code",
    "Scientific Name",
    "S",
    "NP",
    "Type",
    "Size Min",
    "Size Max",
    "Unit",
    "Abundance",
    "Data Quality",
    "Population",
    "Conservation",
    "Isolation",
    "Global",
)


def load_config(path: Path) -> tuple[list[str], set[str]]:
    with path.open(encoding="utf-8") as input_file:
        config = json.load(input_file)

    if not isinstance(config, dict):
        raise ValueError

    urls = config.get("urls")
    needed_animals = config.get("needed_animals")
    if not isinstance(urls, list) or not isinstance(needed_animals, list):
        raise ValueError

    return urls, set(needed_animals)


def get_site_code(url: str) -> str:
    parsed_url = urlsplit(url)
    fragment_query = parsed_url.fragment.partition("?")[2]
    site_codes = parse_qs(parsed_url.query).get("site") or parse_qs(
        fragment_query
    ).get("site")
    if not site_codes or not site_codes[0]:
        raise ValueError
    return site_codes[0]


def get_api_data(
    session: requests.Session,
    endpoint: str,
    params: Mapping[str, str | int],
) -> dict[str, Any]:
    response = session.get(
        f"{API_BASE_URL}/{endpoint}",
        params=params,
        timeout=30,
    )
    response.raise_for_status()

    payload = response.json()
    if not isinstance(payload, dict) or not payload.get("Success"):
        raise RuntimeError
    return payload


def fetch_species(
    session: requests.Session,
    site_code: str,
) -> list[dict[str, Any]]:
    common_params = {"siteCode": site_code, "sdfPublic": "true"}
    release_payload = get_api_data(
        session,
        "GetLastReleaseId",
        common_params,
    )
    release_id = release_payload.get("Data")
    if not isinstance(release_id, int):
        raise ValueError

    data_payload = get_api_data(
        session,
        "GetReleaseDataLegacy",
        {**common_params, "releaseId": release_id},
    )
    data = data_payload.get("Data")
    ecological_information = (
        data.get("EcologicalInformation") if isinstance(data, dict) else None
    )
    species = (
        ecological_information.get("Species")
        if isinstance(ecological_information, dict)
        else None
    )
    if not isinstance(species, list):
        raise ValueError
    return species


def species_row(species: dict[str, Any], zone: str) -> str:
    group = species.get("Group")
    values = (
        zone,
        str(group)[0].upper() if group else "",
        species.get("Code"),
        species.get("SpeciesName"),
        species.get("Sensitive"),
        species.get("NP"),
        species.get("Type"),
        species.get("Min"),
        species.get("Max"),
        species.get("Unit"),
        species.get("Category"),
        species.get("DataQuality"),
        species.get("Population"),
        species.get("Conservation"),
        species.get("Isolation"),
        species.get("Global"),
    )
    return "&".join("" if value is None else str(value) for value in values)


def write_workbook(output_lines: list[str], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError
    sheet.title = "Species"
    header_row = 2

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="1F4E78")

    for row_number, line in enumerate(output_lines, start=header_row + 1):
        values = line.split("&")
        if len(values) != len(HEADERS):
            raise ValueError
        for column, value in enumerate(values, start=1):
            if column in (8, 9) and value:
                value = float(value) if "." in value else int(value)
            sheet.cell(row=row_number, column=column, value=value)

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A{header_row}:P{header_row + len(output_lines)}"
    for column, width in enumerate(
        (14, 10, 10, 28, 6, 6, 8, 10, 10, 8, 12, 13, 12, 14, 10, 8),
        start=1,
    ):
        sheet.column_dimensions[get_column_letter(column)].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    urls, needed_animals = load_config(INPUT_PATH)
    output_lines: list[str] = []

    with requests.Session() as session:
        for url in urls:
            site_code = get_site_code(url)
            species = fetch_species(session, site_code)
            site_lines = [
                species_row(record, site_code)
                for record in species
                if record.get("Code") in needed_animals
            ]
            output_lines.extend(site_lines)
            print(f"{site_code}: found {len(site_lines)} matching rows")

    write_workbook(output_lines, OUTPUT_PATH)
    print(output_lines)
    print(f"Wrote {len(output_lines)} rows to {OUTPUT_PATH}")
